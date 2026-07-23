from __future__ import annotations

from src.models import COURSE_AI_ANALYSIS, COURSE_AI_IMPLEMENTATION
from src.parsers.excel_schedule_parser import group_by_round, parse_excel_schedule


def test_excel_sample_total_candidates_around_40(sample_excel_path):
    with open(sample_excel_path, "rb") as f:
        result = parse_excel_schedule(f, target_year=2026)
    assert 38 <= len(result.events) <= 42


def test_excel_sample_implementation_course_count(sample_excel_path):
    with open(sample_excel_path, "rb") as f:
        result = parse_excel_schedule(f, target_year=2026)
    assert len(result.by_course(COURSE_AI_IMPLEMENTATION)) == 22


def test_excel_sample_analysis_course_count(sample_excel_path):
    with open(sample_excel_path, "rb") as f:
        result = parse_excel_schedule(f, target_year=2026)
    assert len(result.by_course(COURSE_AI_ANALYSIS)) == 20


def test_excel_sample_nine_rounds_per_course(sample_excel_path):
    with open(sample_excel_path, "rb") as f:
        result = parse_excel_schedule(f, target_year=2026)
    impl_rounds = group_by_round(result.by_course(COURSE_AI_IMPLEMENTATION))
    analysis_rounds = group_by_round(result.by_course(COURSE_AI_ANALYSIS))
    assert len(impl_rounds) == 9
    assert len(analysis_rounds) == 9


def test_excel_sample_round1_and_final_are_common(sample_excel_path):
    with open(sample_excel_path, "rb") as f:
        result = parse_excel_schedule(f, target_year=2026)
    impl_rounds = group_by_round(result.by_course(COURSE_AI_IMPLEMENTATION))
    analysis_rounds = group_by_round(result.by_course(COURSE_AI_ANALYSIS))
    assert len(impl_rounds["第1回"]) == len(analysis_rounds["第1回"]) == 2
    assert len(impl_rounds["最終成果発表会"]) == len(analysis_rounds["最終成果発表会"]) == 1


def test_excel_sample_weekday_mismatch_warning(sample_excel_path):
    with open(sample_excel_path, "rb") as f:
        result = parse_excel_schedule(f, target_year=2026)
    assert any("曜日不一致" in w for w in result.warnings)


def test_excel_missing_sheet_falls_back_to_active_sheet(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "別のシート名"
    ws["A1"] = "第1回"
    ws["B1"] = "7/27(月)20:00-21:30"
    ws["C1"] = "山田"
    path = tmp_path / "other_sheet.xlsx"
    wb.save(path)

    with open(path, "rb") as f:
        result = parse_excel_schedule(f, target_year=2026)

    assert len(result.events) >= 1
    assert any("見つからなかった" in w for w in result.warnings)
