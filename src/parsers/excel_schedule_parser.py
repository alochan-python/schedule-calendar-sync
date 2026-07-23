"""ライブ講義スケジュール(Excel)を解析するパーサー。

想定シート構造(A1:E28付近):
    1〜2行目: タイトル・注意書き
    3〜5行目: 第1回の共通講義(AI実装・AI分析コース共通)
    7行目  : コース見出し(AI実装コース / AI分析コース)
    8行目  : 列見出し(日時 / 担当者)
    9行目〜: 第2回〜第8回(コースごとに複数の日時候補があり得る)
    最終行 : 最終成果発表会(共通)

A列の「第○回」が空欄の行は、直前の値を引き継ぐ(同じ回の複数候補)。
「第1回」「最終成果発表会」はAI実装・AI分析の両コースで共通に利用できるようにする。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Optional

import openpyxl

from src.models import (
    COURSE_AI_ANALYSIS,
    COURSE_AI_IMPLEMENTATION,
    SOURCE_TYPE_EXCEL,
    STATUS_CANDIDATE,
    ScheduleEvent,
    make_excel_source_key,
)
from src.parsers.datetime_utils import DateTimeCandidate, combine, find_excel_datetime_candidates

SHEET_NAME = "ライブ講義スケジュール"

_COMMON_ROUND_KEYWORDS = ["第1回", "最終成果発表会", "最終発表", "最終回"]


@dataclass
class ExcelParseResult:
    events: list[ScheduleEvent] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    sheet_name: str = ""

    def by_course(self, course: str) -> list[ScheduleEvent]:
        return [e for e in self.events if e.course == course]


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_common_round(round_name: str) -> bool:
    return any(keyword in round_name for keyword in _COMMON_ROUND_KEYWORDS)


def _extract_column_candidates(
    cell_text: str, instructor_text: str, year: int
) -> list[tuple[DateTimeCandidate, str]]:
    candidates = find_excel_datetime_candidates(cell_text, year)
    if not candidates:
        return []
    instructor_lines = [line.strip() for line in (instructor_text or "").split("\n") if line.strip()]
    results = []
    for i, candidate in enumerate(candidates):
        if instructor_lines:
            instructor = instructor_lines[i] if i < len(instructor_lines) else instructor_lines[-1]
        else:
            instructor = ""
        results.append((candidate, instructor))
    return results


def _make_event(
    course: str,
    round_name: str,
    candidate: DateTimeCandidate,
    instructor: str,
    source_file: str,
) -> ScheduleEvent:
    start_dt = combine(candidate.date, candidate.start_time)
    end_dt = combine(candidate.date, candidate.end_time)
    description_lines = [
        f"講師: {instructor or '未定'}",
        f"コース: {course}",
        f"回: {round_name}",
        "取込元: Excel(ライブ講義スケジュール)",
    ]
    return ScheduleEvent(
        source_type=SOURCE_TYPE_EXCEL,
        source_file=source_file,
        source_key=make_excel_source_key(course, round_name, start_dt),
        course=course,
        session_name=round_name,
        title=f"{course}｜{round_name}",
        start_datetime=start_dt,
        end_datetime=end_dt,
        location="オンライン",
        instructor=instructor,
        description="\n".join(description_lines),
        status=STATUS_CANDIDATE,
        weekday=candidate.weekday_written,
    )


def parse_excel_schedule(
    file_like,
    target_year: int,
    source_file: str = "excel_schedule.xlsx",
) -> ExcelParseResult:
    """アップロードされたExcelファイル(ファイルオブジェクトまたはパス)を解析する。"""
    result = ExcelParseResult()

    workbook = openpyxl.load_workbook(file_like, data_only=True)
    if SHEET_NAME in workbook.sheetnames:
        sheet = workbook[SHEET_NAME]
        result.sheet_name = SHEET_NAME
    else:
        sheet = workbook.active
        result.sheet_name = sheet.title
        result.warnings.append(
            f"シート「{SHEET_NAME}」が見つからなかったため、先頭シート「{sheet.title}」を使用します。"
        )

    current_round: Optional[str] = None

    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_col=5), start=1):
        cells = list(row) + [None] * (5 - len(row))
        a_text = _cell_text(cells[0].value if cells[0] else None)
        b_text = _cell_text(cells[1].value if cells[1] else None)
        c_text = _cell_text(cells[2].value if cells[2] else None)
        d_text = _cell_text(cells[3].value if cells[3] else None)
        e_text = _cell_text(cells[4].value if cells[4] else None)

        if a_text:
            current_round = a_text

        if current_round is None:
            continue  # タイトル・注意書き行はスキップ

        is_common = _is_common_round(current_round)

        impl_candidates = _extract_column_candidates(b_text, c_text, target_year)
        analysis_candidates = _extract_column_candidates(d_text, e_text, target_year)

        if is_common:
            if impl_candidates and not analysis_candidates:
                analysis_candidates = impl_candidates
            elif analysis_candidates and not impl_candidates:
                impl_candidates = analysis_candidates

        for candidate, instructor in impl_candidates:
            if candidate.weekday_mismatch:
                result.warnings.append(
                    f"{row_idx}行目(AI実装コース・{current_round}): 曜日不一致の可能性があります"
                    f"(記載:{candidate.weekday_written} 実際:{candidate.weekday_actual})"
                )
            result.events.append(
                _make_event(COURSE_AI_IMPLEMENTATION, current_round, candidate, instructor, source_file)
            )

        for candidate, instructor in analysis_candidates:
            if candidate.weekday_mismatch:
                result.warnings.append(
                    f"{row_idx}行目(AI分析コース・{current_round}): 曜日不一致の可能性があります"
                    f"(記載:{candidate.weekday_written} 実際:{candidate.weekday_actual})"
                )
            result.events.append(
                _make_event(COURSE_AI_ANALYSIS, current_round, candidate, instructor, source_file)
            )

    return result


def group_by_round(events: list[ScheduleEvent]) -> dict[str, list[ScheduleEvent]]:
    """回(session_name)ごとに候補をまとめる。表示順は最初に出現した順を維持する。"""
    grouped: dict[str, list[ScheduleEvent]] = {}
    for event in events:
        grouped.setdefault(event.session_name, []).append(event)
    return grouped
